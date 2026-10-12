"""Exporta el CSV sintético a un Excel legible con diccionario y advertencias."""
import argparse
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

HEADERS=['idx','sentence','label','source','split']
def export(csv_path, xlsx_path):
    csv_path=Path(csv_path); xlsx_path=Path(xlsx_path)
    with csv_path.open(encoding='utf-8',newline='') as f:
        data=list(csv.DictReader(f))
    assert data and list(data[0])==HEADERS
    wb=Workbook();ws=wb.active;ws.title='Dataset'
    ws.append(HEADERS)
    for item in data:ws.append([item['idx'],item['sentence'],int(item['label']),item['source'],item['split']])
    ws.freeze_panes='B2';ws.sheet_view.showGridLines=False
    for col,w in {'A':18,'B':94,'C':14,'D':24,'E':18}.items():ws.column_dimensions[col].width=w
    ws.row_dimensions[1].height=28
    for cell in ws[1]:
        cell.fill=PatternFill('solid',fgColor='123754');cell.font=Font(name='Aptos',bold=True,color='FFFFFF',size=11)
        cell.alignment=Alignment(vertical='center')
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font=Font(name='Aptos',size=10,color='00804D' if cell.column<=5 else '333333')
            cell.alignment=Alignment(vertical='center',wrap_text=False)
    tab=Table(displayName='OpinionesSinteticas',ref=f'A1:E{len(data)+1}')
    tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showFirstColumn=False,showLastColumn=False,showRowStripes=True,showColumnStripes=False)
    ws.add_table(tab)
    w=wb.create_sheet('Diccionario');w.sheet_view.showGridLines=False
    for r in [['Campo','Tipo','Significado'],
              ['idx','Texto','Clave única SYN-label-NNN'],
              ['sentence','Texto','Frase corta original en inglés'],
              ['label','Entero','0=negativo, 1=positivo'],
              ['source','Texto','synthetic_manual: contenido creado para la actividad'],
              ['split','Texto','smoke_test: no se usa para selección de modelos'],
              [],['Advertencia','La evaluación principal usa GLUE SST-2 validation (872 filas) y se descarga desde el notebook.'],
              ['Fuente SST-2','https://huggingface.co/datasets/nyu-mll/glue'],
              ['Generación','python src/generate_synthetic.py']]:w.append(r)
    for cell in w[1]:cell.font=Font(name='Aptos',color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='123754')
    w.column_dimensions['A'].width=25;w.column_dimensions['B'].width=74;w.column_dimensions['C'].width=63
    for row in w.iter_rows(min_row=2):
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
    xlsx_path.parent.mkdir(parents=True,exist_ok=True);wb.save(xlsx_path)
    print(f'Excel generado: {xlsx_path} ({len(data)} registros)')

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--csv',default='data/synthetic_smoke_test.csv');p.add_argument('--out',default='data/dataset_sintetico.xlsx')
    a=p.parse_args();export(a.csv,a.out)
